from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from datetime import date

candidats = [
    {"nom": "Pauline Auche Harreau",    "telephone": "+33 6 50 51 29 26"},
    {"nom": "Lucas Hunot",               "telephone": "07 83 65 01 26"},
    {"nom": "Coralie Fleuriot",          "telephone": "06 27 20 59 35"},
    {"nom": "Pierre-Laurent Forestier",  "telephone": "07 60 59 56 84"},
    {"nom": "Alexandre Gascon",          "telephone": "06 27 16 41 42"},
    {"nom": "Chloé Ledoyen",             "telephone": "06 21 13 47 43"},
]

doc = Document()

# Marges
for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# Titre
titre = doc.add_heading("CV à rappeler", level=1)
titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
titre.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

# Sous-titre date
sous = doc.add_paragraph(f"Liste établie le {date.today().strftime('%d/%m/%Y')}")
sous.alignment = WD_ALIGN_PARAGRAPH.CENTER
sous.runs[0].font.size = Pt(10)
sous.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

doc.add_paragraph("")

# Tableau
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# En-têtes
entetes = ["#", "Nom", "Téléphone", "Rappelé ?"]
for i, e in enumerate(entetes):
    cell = table.rows[0].cells[i]
    cell.text = e
    run = cell.paragraphs[0].runs[0]
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "1F497D")
    tcPr.append(shd)

# Lignes candidats
for idx, c in enumerate(candidats):
    row = table.add_row()
    row.cells[0].text = str(idx + 1)
    row.cells[1].text = c["nom"]
    row.cells[2].text = c["telephone"]
    row.cells[3].text = "☐"
    for cell in row.cells:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    # Alternance couleur
    if idx % 2 == 0:
        for cell in row.cells:
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "DCE6F1")
            tcPr.append(shd)

# Largeurs colonnes
from docx.oxml.ns import qn
widths = [Cm(1), Cm(6), Cm(4), Cm(3)]
for row in table.rows:
    for i, cell in enumerate(row.cells):
        cell.width = widths[i]

doc.add_paragraph("")
note = doc.add_paragraph("Les CV PDF correspondants sont disponibles dans le dossier : C:\\Users\\gaeta\\Desktop\\CV")
note.runs[0].font.size = Pt(9)
note.runs[0].font.italic = True
note.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

doc.save("/home/user/app_soldes/CV_a_rappeler.docx")
print("Document généré : CV_a_rappeler.docx")
